# -*- coding: utf-8 -*-
"""验证 docx 内容并导出 PDF"""
import os
import sys

from docx import Document

docx_path = sys.argv[1]
pdf_path = sys.argv[2]

d = Document(docx_path)
print("paragraphs:", len(d.paragraphs))
print("tables    :", len(d.tables))
heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading") and p.text.strip()]
print("headings  :", len(heads))
for h in heads[:14]:
    print("   -", h[:70])

# 导出 PDF
try:
    import win32com.client

    w = win32com.client.gencache.EnsureDispatch("Word.Application")
    w.Visible = False
    w.DisplayAlerts = 0
    doc = w.Documents.Open(os.path.abspath(docx_path))
    doc.Repaginate()
    pages = doc.ComputeStatistics(2)
    words = doc.ComputeStatistics(0)
    chars = doc.ComputeStatistics(3)
    print(f"pages={pages}  words={words}  chars={chars}")
    doc.ExportAsFixedFormat(OutputFileName=os.path.abspath(pdf_path), ExportFormat=17)
    doc.Close(False)
    w.Quit()
    print("pdf saved:", pdf_path)
except Exception as e:
    print("PDF export failed:", e)
